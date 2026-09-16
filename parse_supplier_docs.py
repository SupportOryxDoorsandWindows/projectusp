#!/usr/bin/env python3
"""
Oryx Doors & Windows - messy supplier document parser.

build_kb.py reads one specific spreadsheet whose sheet names, rows and
columns are hardcoded (see BLOCKS). That works for the maintained "Copy of
Slider USP.xlsx" file, but supplier product sheets arrive as PDFs and Word
documents with no fixed layout: column order varies, headers are worded
differently from one supplier to the next, and some pages are scanned
images with no extractable text at all.

This module reads those PDFs and DOCX files by matching table headers
against known field names (fuzzy, not positional), instead of assuming a
fixed cell layout. It never guesses at data it cannot identify:

  - A recognised field with an empty cell is recorded as `None` ("not
    recorded"), matching the tri-state convention build_kb.py already uses.
  - A column whose header cannot be matched to a known field is kept under
    its own (cleaned) header text rather than being silently dropped.
  - A page or table that yields no usable text is reported as unreadable
    (e.g. a scanned image with no text layer) instead of being merged in
    as if it were empty data.
  - Two different systems that would produce the same id (same name, seen
    twice) are reported as a conflict rather than one silently overwriting
    the other.

Usage
-----
    python3 parse_supplier_docs.py supplier_sheet.pdf catalog.docx
    python3 parse_supplier_docs.py supplier_sheet.pdf --merge-into data/kb.json

With --merge-into, parsed systems are added to (or replace, by id) the
`systems` list of an existing kb.json/kb.js pair, tagged with their source
file so they can be told apart from the spreadsheet-derived ones. Without
it, results are printed as JSON plus a human-readable review report on
stderr, and nothing is written.
"""

import argparse
import json
import os
import re
import sys
import unicodedata

# --------------------------------------------------------------------------
# Text cleanup
# --------------------------------------------------------------------------
# Suppliers export from many different tools, so the same "no data" concept
# shows up as several different tokens. These two sets are deliberately
# separate: both display as blank, but only EXPLICIT_NA_TOKENS should make a
# tri-state option field False ("not available"); BLANK_TOKENS ("pending"
# markers) must land on None ("not recorded"), same as a genuinely empty
# cell - a "TBD" is not a confirmed no.
EXPLICIT_NA_TOKENS = {"n/a", "na", "n.a.", "n.a", "not applicable", "not available"}
BLANK_TOKENS = {"tbd", "tbc"}

# Characters that render as a hyphen/dash to a human but are not "-" (U+002D).
DASH_CHARS = "‐‑‒–—―−"

UNIT_RE = re.compile(r"\b(?:mm|cm|m2|m²|sq\.?\s*m\.?|sqm)\b\.?", re.I)


def _normalize(value):
    """NFKC-normalise, strip control characters and non-'-' dashes/nbsp,
    collapse whitespace. Shared by clean_text() and is_na() so both agree
    on what a cell's text actually is."""
    if value is None:
        return ""
    s = unicodedata.normalize("NFKC", str(value))
    s = "".join(ch for ch in s if unicodedata.category(ch)[0] not in ("C",) or ch in "\n\t")
    for d in DASH_CHARS:
        s = s.replace(d, "-")
    s = s.replace(" ", " ")
    return re.sub(r"\s+", " ", s).strip()


def clean_text(value):
    """Collapse whitespace and normalise dashes/N/A tokens.

    Returns None for a genuinely empty, N/A, or "pending" cell, so callers
    can apply the "not recorded" (None) vs "not available" (False, via
    is_na()) distinction the rest of the knowledge base uses.
    """
    s = _normalize(value)
    if s == "":
        return None
    low = s.lower()
    if low in EXPLICIT_NA_TOKENS or low in BLANK_TOKENS or set(s) <= set("- "):
        return None
    return s


def is_na(value):
    """True only for an explicit "not available" cell - not a blank, and
    not a "pending" marker like TBD/TBC, which means not yet recorded."""
    s = _normalize(value)
    if s == "":
        return False
    low = s.lower()
    return low in EXPLICIT_NA_TOKENS or set(s) <= set("- ")


def normalize_header(value):
    """Fold a header down to bare words for fuzzy matching."""
    s = clean_text(value) or ""
    s = s.lower()
    s = re.sub(r"\(.*?\)", " ", s)  # drop parenthetical units, e.g. "(mm)"
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _strip_units(t):
    return re.sub(r"\s+", " ", UNIT_RE.sub(" ", t)).strip()


def parse_range(text):
    """'216-1000' / '216 to 1000' / '1500' / '400-1,800 mm' -> (min, max)."""
    cleaned = clean_text(text)
    if cleaned is None:
        return (None, None)
    t = _strip_units(cleaned.replace(",", ""))
    m = re.match(r"^(\d+(?:\.\d+)?)\s*(?:-|to|~)\s*(\d+(?:\.\d+)?)$", t, re.I)
    if m:
        return (float(m.group(1)), float(m.group(2)))
    m = re.match(r"^(\d+(?:\.\d+)?)$", t)
    if m:
        return (None, float(m.group(1)))
    return (None, None)


def parse_single_number(text):
    """A cell that should hold exactly one number, e.g. an explicit 'Min
    Width' or 'Max Width' column (as opposed to a combined range column)."""
    cleaned = clean_text(text)
    if cleaned is None:
        return None
    t = _strip_units(cleaned.replace(",", ""))
    m = re.match(r"^(\d+(?:\.\d+)?)$", t)
    return float(m.group(1)) if m else None


def slugify(name):
    s = normalize_header(name).replace(" ", "-")
    return re.sub(r"-+", "-", s).strip("-") or "unnamed"


def leaves_in_config(cfg):
    """'Fixed + Sliding + Sliding + Fixed' -> (4 leaves, 2 operable).
    Mirrors build_kb.py's convention for the same column."""
    parts = [p.strip().lower() for p in cfg.split("+") if p.strip()]
    return len(parts), sum(1 for p in parts if p.startswith("slid"))


# --------------------------------------------------------------------------
# Field matching - header text -> canonical schema field
# --------------------------------------------------------------------------
# Aliases here are safe to match as a *word* anywhere in a header (e.g.
# "width" matches both "Width" and "Max Width"). Generic single words that
# are ambiguous out of context (a bare "product", "system", "type") are
# deliberately NOT here - they live in FIELD_EXACT_ALIASES instead, and only
# match a header that is *exactly* that word, so "Product Width" resolves
# to the width field rather than being stolen by "product".
FIELD_ALIASES = {
    "name": ["system name", "product name", "name"],
    "family": ["product type", "system type", "family"],
    "sash_w_min": ["min width", "minimum width"],
    "sash_w_max": ["max width", "maximum width"],
    "sash_w_range": ["width", "sash width", "panel width", "w"],
    "sash_h_min": ["min height", "minimum height"],
    "sash_h_max": ["max height", "maximum height"],
    "sash_h_range": ["height", "sash height", "panel height", "h"],
    "sash_sqm_max": ["max area", "max sqm", "sash area", "area", "sqm"],
    "glass": ["glass thickness", "glazing", "glass"],
    "automation": ["automation"],
    "locking": ["locking options", "locking", "lock"],
    "configuration": ["leaf configuration", "panel configuration", "configuration", "config"],
}
FIELD_EXACT_ALIASES = {
    "name": ["product", "system", "series", "model"],
    "family": ["type", "category"],
}

# Known option sub-types, matched before falling back to a generic bucket.
OPTION_ALIASES = {
    "threshold": {
        "floor integrated": ["floor integrated", "floor integrated marble or wood"],
        "floor flushed": ["floor flushed", "flush floor"],
        "on top of ffl": ["on top of ffl", "above ffl"],
        "stepped floor": ["stepped floor"],
    },
    "drainage": {
        "visible drainage": ["visible drainage"],
        "concealed drainage": ["concealed drainage", "hidden drainage"],
    },
    "sightline": {
        "horizontal sightline": ["horizontal sightline"],
        "vertical sightline": ["vertical sightline"],
    },
}
OPTION_KIND_WORDS = {"threshold": "threshold", "drainage": "drainage", "sightline": "sightline"}


def _match_alias(header_norm, alias_map):
    best_canonical, best_len = None, -1
    for canonical, aliases in alias_map.items():
        for alias in aliases:
            if header_norm == alias:
                return canonical
            # Word-boundary match so short aliases ("w", "h") only match a
            # whole word (header "max w") and don't fire on substrings
            # inside an unrelated word. Longest alias wins so a specific
            # phrase ("system type") beats a shorter unrelated one that
            # also happens to match ("system" for name).
            if re.search(r"(?:^|\s)%s(?:$|\s)" % re.escape(alias), header_norm):
                if len(alias) > best_len:
                    best_canonical, best_len = canonical, len(alias)
    return best_canonical


def match_field(header):
    """Classify one table header. Returns (kind, key) where kind is one of
    'field', 'threshold', 'drainage', 'sightline', 'unknown'."""
    norm = normalize_header(header)
    if not norm:
        return ("unknown", None)

    # Option words (threshold/drainage/sightline) are the most specific,
    # unambiguous signal available - check them before the generic field
    # aliases so e.g. "Threshold Type" is captured as a threshold column,
    # not misread as the system's family via the word "type".
    for kind, subtypes in OPTION_ALIASES.items():
        sub = _match_alias(norm, subtypes)
        if sub:
            return (kind, sub)
    for kind, word in OPTION_KIND_WORDS.items():
        if re.search(r"(?:^|\s)%s(?:$|\s)" % word, norm):
            return (kind, header.strip())

    for canonical, words in FIELD_EXACT_ALIASES.items():
        if norm in words:
            return ("field", canonical)

    field = _match_alias(norm, FIELD_ALIASES)
    if field:
        return ("field", field)

    return ("unknown", header.strip())


# --------------------------------------------------------------------------
# Table extraction
# --------------------------------------------------------------------------
def _collapse_merged_docx_row(row):
    """python-docx repeats the same cell text once per grid column for a
    horizontally merged cell (e.g. a title spanning 4 columns comes back as
    4 identical cells). Collapse runs of identical non-blank text down to
    their first occurrence so header-row scoring doesn't see phantom
    repeated columns."""
    out = list(row)
    i = 0
    while i < len(out):
        j = i + 1
        while j < len(out) and out[i] != "" and out[j] == out[i]:
            out[j] = ""
            j += 1
        i = j
    return out


def extract_tables_from_pdf(path):
    """Yield (page_number, table_rows, unreadable_reason) for every table
    pdfplumber can find. unreadable_reason is None when table_rows is
    usable data; otherwise table_rows is None and the reason distinguishes
    a genuinely blank/scanned page from one that has running text pdfplumber
    could not organise into a table - both need a human to look at them,
    but for different reasons.
    """
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            page_tables = page.extract_tables() or []
            if page_tables:
                for t in page_tables:
                    yield (i, t, None)
                continue
            text = page.extract_text() or ""
            if text.strip():
                yield (i, None, "has text but no table structure was detected - content not parsed")
            else:
                yield (i, None, "no extractable text - likely a scanned image; needs OCR or manual entry")


def extract_tables_from_docx(path):
    import docx

    document = docx.Document(path)
    if not document.tables:
        text = "\n".join(p.text for p in document.paragraphs)
        if text.strip():
            yield (None, None, "document has text but no tables - content not parsed")
        else:
            yield (None, None, "no tables and no text found")
        return
    for idx, table in enumerate(document.tables, start=1):
        rows = [_collapse_merged_docx_row([cell.text for cell in row.cells]) for row in table.rows]
        yield (idx, rows, None)


# --------------------------------------------------------------------------
# Table -> system records
# --------------------------------------------------------------------------
def table_to_records(rows, source_label):
    """Turn one table's rows into system dicts plus a review report.

    The header row is whichever of the first few rows matches the most
    known fields (preferring one that includes a recognisable name column) -
    some exports add a title row above the real header.
    """
    report = {"unmatched_headers": [], "empty_pages": [], "no_name": 0}
    if not rows or len(rows) < 2:
        return [], report

    header_row_idx, header_map = _find_header_row(rows)
    if header_map is None:
        report["unmatched_headers"].append(
            "%s: no row matched a recognisable header, table skipped" % source_label
        )
        return [], report

    records = []
    for r in rows[header_row_idx + 1:]:
        if all(clean_text(c) is None for c in r):
            continue  # blank spacer row
        record = _row_to_record(r, header_map)
        if record.get("name") is None:
            report["no_name"] += 1
            continue
        record["_source"] = source_label
        records.append(record)

    for col_idx, header in enumerate(rows[header_row_idx]):
        kind, key = match_field(header)
        if kind == "unknown" and clean_text(header) is not None:
            report["unmatched_headers"].append(
                "%s: column %r not recognised - kept as free text" % (source_label, clean_text(header))
            )
    return records, report


def _find_header_row(rows, scan_limit=3):
    best_idx, best_map, best_score, best_has_name = None, None, 0, False
    for idx in range(min(scan_limit, len(rows))):
        header_map = {}
        score = 0
        has_name = False
        for col_idx, header in enumerate(rows[idx]):
            kind, key = match_field(header)
            header_map[col_idx] = (kind, key)
            if kind != "unknown":
                score += 1
            if kind == "field" and key == "name":
                has_name = True
        if score == 0:
            continue
        # A row with a recognisable name column is preferred outright over
        # one without, even at a lower raw score - a merged title row
        # repeating one word rarely names an actual "name" column.
        if best_map is None or (has_name, score) > (best_has_name, best_score):
            best_idx, best_map, best_score, best_has_name = idx, header_map, score, has_name
    if best_map is None:
        return 0, None
    return best_idx, best_map


def _row_to_record(row, header_map):
    record = {
        "name": None, "family": None,
        "sash_w_min": None, "sash_w_max": None,
        "sash_h_min": None, "sash_h_max": None,
        "sash_sqm_max": None, "glass": None,
        "automation": None, "locking": None, "configuration": None,
        "thresholds": {}, "drainage": {}, "sightlines": {}, "notes": {},
    }
    for col_idx, cell in enumerate(row):
        if col_idx not in header_map:
            continue
        kind, key = header_map[col_idx]
        value = clean_text(cell)
        na = is_na(cell)

        if kind == "field":
            if key == "name":
                record["name"] = value
            elif key == "family":
                record["family"] = value
            elif key == "sash_w_min":
                v = parse_single_number(cell)
                if v is not None:
                    record["sash_w_min"] = v
            elif key == "sash_w_max":
                v = parse_single_number(cell)
                if v is not None:
                    record["sash_w_max"] = v
            elif key == "sash_w_range":
                lo, hi = parse_range(cell)
                if lo is not None:
                    record["sash_w_min"] = lo
                if hi is not None:
                    record["sash_w_max"] = hi
            elif key == "sash_h_min":
                v = parse_single_number(cell)
                if v is not None:
                    record["sash_h_min"] = v
            elif key == "sash_h_max":
                v = parse_single_number(cell)
                if v is not None:
                    record["sash_h_max"] = v
            elif key == "sash_h_range":
                lo, hi = parse_range(cell)
                if lo is not None:
                    record["sash_h_min"] = lo
                if hi is not None:
                    record["sash_h_max"] = hi
            elif key == "sash_sqm_max":
                _, v = parse_range(cell)
                if v is not None:
                    record["sash_sqm_max"] = v
            else:
                record[key] = value
        elif kind in ("threshold", "drainage", "sightline"):
            bucket = record["thresholds"] if kind == "threshold" else (
                record["drainage"] if kind == "drainage" else record["sightlines"])
            bucket[key] = False if na else (True if value is not None else None)
        elif kind == "unknown" and key:
            if value is not None:
                record["notes"][key] = value
    return record


# --------------------------------------------------------------------------
# Top-level parsing
# --------------------------------------------------------------------------
def parse_document(path):
    """Parse one supplier file. Never raises for a bad/unreadable file -
    problems are collected into the returned report so one broken file
    does not stop the rest of a batch from being processed."""
    records, report = [], {"unmatched_headers": [], "empty_pages": [], "no_name": 0, "errors": []}
    ext = os.path.splitext(path)[1].lower()

    try:
        if ext == ".pdf":
            source_tables = extract_tables_from_pdf(path)
        elif ext == ".docx":
            source_tables = extract_tables_from_docx(path)
        else:
            report["errors"].append("%s: unsupported file type %r" % (path, ext))
            return records, report

        for page_or_table_num, rows, reason in source_tables:
            if rows is None:
                report["empty_pages"].append("%s (page/table %s): %s" % (path, page_or_table_num, reason))
                continue
            label = "%s#%s" % (os.path.basename(path), page_or_table_num)
            recs, sub_report = table_to_records(rows, label)
            records.extend(recs)
            report["unmatched_headers"].extend(sub_report["unmatched_headers"])
            report["no_name"] += sub_report["no_name"]
    except Exception as exc:  # noqa: BLE001 - a bad file must not kill the batch
        report["errors"].append("%s: failed to parse (%s: %s)" % (path, type(exc).__name__, exc))

    return records, report


def record_to_system(record):
    """A supplier system's id is derived purely from its (slugified) name -
    stable across reruns of the same input, so --merge-into can replace the
    same system in place instead of accumulating duplicates. Two distinct
    records that happen to produce the same id are a genuine naming
    collision the caller must report and resolve, not something this
    function should silently paper over with a suffix."""
    sid = "supplier:" + slugify(record["name"])

    configs = []
    any_config = False
    cfg = record["configuration"]
    if cfg:
        if cfg.lower().startswith("any"):
            any_config = True
            configs.append(dict(label="Any configuration", leaves=None, operable=None, tracks=None))
        else:
            n, op = leaves_in_config(cfg)
            configs.append(dict(label=cfg, leaves=n, operable=op, tracks=None))

    return dict(
        id=sid, name=record["name"], family=record["family"],
        sash_w_min=record["sash_w_min"], sash_w_max=record["sash_w_max"],
        sash_h_min=record["sash_h_min"], sash_h_max=record["sash_h_max"],
        sash_sqm_max=record["sash_sqm_max"],
        glass=record["glass"], automation=record["automation"], locking=record["locking"],
        configs=configs, any_config=any_config, tracks=[],
        thresholds=record["thresholds"], drainage=record["drainage"],
        sightlines=record["sightlines"], drawings=[],
        source_notes=record["notes"], source=record["_source"],
    )


def _dedupe_by_id(records, report):
    """Convert records to systems, reporting (and dropping) same-batch id
    collisions instead of letting a later record silently replace an
    earlier, different one."""
    systems, seen = [], {}
    for r in records:
        system = record_to_system(r)
        sid = system["id"]
        if sid in seen:
            report.setdefault("id_conflicts", []).append(
                "%r (id %s) seen in both %s and %s - kept the first, second skipped"
                % (r["name"], sid, seen[sid], r["_source"])
            )
            continue
        seen[sid] = r["_source"]
        systems.append(system)
    return systems


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------
def print_report(report):
    if report.get("errors"):
        print("Files that could not be read at all:", file=sys.stderr)
        for e in report["errors"]:
            print("  - %s" % e, file=sys.stderr)
    if report.get("empty_pages"):
        print("Unreadable pages/tables (need a human to look at the source file):", file=sys.stderr)
        for e in report["empty_pages"]:
            print("  - %s" % e, file=sys.stderr)
    if report.get("unmatched_headers"):
        print("Columns kept as free text (header not recognised):", file=sys.stderr)
        for e in report["unmatched_headers"]:
            print("  - %s" % e, file=sys.stderr)
    if report.get("id_conflicts"):
        print("Naming conflicts (two systems would get the same id):", file=sys.stderr)
        for e in report["id_conflicts"]:
            print("  - %s" % e, file=sys.stderr)
    if report.get("no_name"):
        print("Rows skipped for having no identifiable system name: %d" % report["no_name"],
              file=sys.stderr)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="+", help="Supplier PDF and/or DOCX files")
    ap.add_argument("--merge-into", metavar="KB_JSON",
                     help="Merge parsed systems into this kb.json (and its sibling kb.js) by id")
    args = ap.parse_args()

    all_records = []
    combined_report = {"unmatched_headers": [], "empty_pages": [], "no_name": 0, "errors": []}
    for path in args.files:
        if not os.path.exists(path):
            combined_report["errors"].append("%s: file not found" % path)
            continue
        records, report = parse_document(path)
        all_records.extend(records)
        for k in combined_report:
            if isinstance(combined_report[k], list):
                combined_report[k].extend(report.get(k, []))
            else:
                combined_report[k] += report.get(k, 0)

    if args.merge_into:
        merge_into_kb(all_records, args.merge_into, combined_report)
    else:
        systems = _dedupe_by_id(all_records, combined_report)
        print_report(combined_report)
        print("\nParsed %d system record(s) from %d file(s)." % (len(systems), len(args.files)),
              file=sys.stderr)
        json.dump(dict(systems=systems), sys.stdout, indent=1)
        print()


def merge_into_kb(records, kb_path, report):
    with open(kb_path) as f:
        kb = json.load(f)

    by_id = {s["id"]: i for i, s in enumerate(kb["systems"])}
    systems = _dedupe_by_id(records, report)

    added, replaced = 0, 0
    for system in systems:
        sid = system["id"]
        if sid in by_id:
            kb["systems"][by_id[sid]] = system
            replaced += 1
        else:
            kb["systems"].append(system)
            by_id[sid] = len(kb["systems"]) - 1
            added += 1

    print_report(report)
    print("\nParsed %d system record(s) from source file(s)." % len(systems), file=sys.stderr)

    with open(kb_path, "w") as f:
        json.dump(kb, f, indent=1)
    js_path = os.path.splitext(kb_path)[0] + ".js"
    if os.path.exists(js_path):
        with open(js_path, "w") as f:
            f.write("window.ORYX_KB = ")
            json.dump(kb, f, indent=1)
            f.write(";\n")

    print("Merged into %s: %d added, %d replaced." % (kb_path, added, replaced), file=sys.stderr)


if __name__ == "__main__":
    main()
